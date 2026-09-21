#!/usr/bin/env python3
"""
Import student data from Repartition-1.xlsx into the public JSON dataset.

  python3 scripts/import_students.py [path/to/Repartition-1.xlsx]
  python3 scripts/import_students.py --demo     # generate the clearly-flagged demo dataset

Rules (per project spec):
  - Primary sheet: "Résultat consolidé"
  - "À vérifier" sheet is IGNORED for the public dataset
  - Only public fields are exported: name, original class, cycle, group, Kid's Book
  - Score / confidence / remarks / ranking are NEVER written to the dataset
  - Empty rows, title rows and repeated header rows are removed
  - Duplicate records (same normalized name + class) are removed
  - Every exported record must have: name, class, group, book
  - "Manuel attribué" is authoritative when present; otherwise the documented
    cycle/group -> book mapping is used as a fallback and flagged in the report
"""

import hashlib
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA_OUT = ROOT / "server" / "data" / "students.json"

SHEET_NAME = "Résultat consolidé"
IGNORED_SHEETS = {"à vérifier", "a verifier"}

# Documented fallback mapping (spec §13). "Manuel attribué" always wins when present.
BOOK_MAPPING = {
    (1, 1): "Kid's Box 1",
    (1, 2): "Kid's Box 2",
    (1, 3): "Kid's Box 2",
    (2, 1): "Kid's Box 3",
    (2, 2): "Kid's Box 3",
    (2, 3): "Kid's Box 4",
    (3, 1): "Kid's Box 5",
    (3, 2): "Kid's Box 5",
    (3, 3): "Kid's Box 6",
}


def norm(s):
    """Normalize a string for matching / dedupe (accent- and case-insensitive)."""
    if s is None:
        return ""
    s = str(s).strip()
    s = unicodedata.normalize("NFKC", s).casefold()
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^0-9a-z]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def cell(row, idx):
    """Safely read a cell from an openpyxl row."""
    if idx < 0 or idx >= len(row):
        return ""
    v = row[idx]
    if v is None:
        return ""
    return str(v).strip()


def detect_header(rows):
    """Find the header row and column indexes. Returns (header_idx, cols) or None."""
    def find(header_cells):
        cols = {"name": None, "class": None, "cycle": None, "group": None, "book": None}
        for i, raw in enumerate(header_cells):
            h = norm(raw)
            if cols["name"] is None and ("nom" in h or "pr" in h or "name" in h):
                cols["name"] = i
            elif cols["class"] is None and ("classe" in h or "class" in h or "origine" in h):
                cols["class"] = i
            elif cols["cycle"] is None and ("cycle" in h or "niveau" in h):
                cols["cycle"] = i
            elif cols["group"] is None and ("groupe" in h or "group" in h):
                cols["group"] = i
            elif cols["book"] is None and ("manuel" in h or "book" in h or "manipule" in h):
                cols["book"] = i
        return cols

    for idx, row in enumerate(rows[:15]):
        cols = find([cell(row, i) for i in range(len(row))])
        if cols["name"] is not None and cols["group"] is not None:
            return idx, cols
    return None


def parse_cycle(v):
    n = norm(v)
    m = re.search(r"(\d)", n)
    if not m:
        return None
    c = int(m.group(1))
    return c if c in (1, 2, 3) else None


def parse_group(v):
    n = norm(v)
    m = re.search(r"(\d)", n)
    if not m:
        return None
    g = int(m.group(1))
    return g if 1 <= g <= 9 else None


def parse_book(v):
    """Normalize a 'Manuel attribué' value to 'Kid's Box N' (or None)."""
    n = norm(v)
    if not n:
        return None
    m = re.search(r"(\d)", n)
    if not m:
        return None
    return f"Kid's Box {int(m.group(1))}"


def record_id(name, cls, group):
    h = hashlib.sha256(f"{norm(name)}|{norm(cls)}|{group}".encode("utf-8")).hexdigest()
    return h[:12]


def build_record(name, cls, cycle_raw, group_raw, book_raw, source, warnings):
    name = name.replace("\n", " ").strip()
    name = re.sub(r"\s{2,}", " ", name)
    cls = cls.strip()
    cycle = parse_cycle(cycle_raw)
    group = parse_group(group_raw)

    if not name or len(name) < 3:
        return None
    if not cls:
        warnings.append(f"Skipped (no class): {name!r}")
        return None
    if group is None:
        warnings.append(f"Skipped (no group): {name!r} ({cls})")
        return None

    book = parse_book(book_raw)
    book_source = "database"
    if book is None:
        if cycle is not None:
            book = BOOK_MAPPING.get((cycle, group))
            book_source = "mapping"
            warnings.append(f"{name!r} ({cls}, G{group}): 'Manuel attribué' missing -> mapping fallback")
        else:
            warnings.append(f"Skipped (no book, no cycle): {name!r} ({cls})")
            return None
    if book is None:
        warnings.append(f"Skipped (unknown book for C{cycle}/G{group}): {name!r} ({cls})")
        return None

    return {
        "id": record_id(name, cls, group),
        "name": name,
        "className": cls,
        "cycle": cycle,
        "group": group,
        "book": book,
        "bookSource": book_source,
    }


def import_workbook(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = [str(s).strip() for s in wb.sheetnames]
    target = None
    for s in sheet_names:
        if norm(s) == norm(SHEET_NAME):
            target = s
            break
    if target is None:
        # tolerate slight naming differences
        for s in sheet_names:
            if "consolid" in norm(s):
                target = s
                break
    if target is None:
        raise SystemExit(f"Sheet {SHEET_NAME!r} not found. Sheets: {sheet_names}")

    ws = wb[target]
    rows = list(ws.iter_rows())
    found = detect_header(rows)
    if found is None:
        raise SystemExit(f"Could not detect the header row in sheet {target!r}")
    header_idx, cols = found

    warnings = []
    students = []
    seen = set()
    skipped_header = 0

    for row in rows[header_idx + 1 :]:
        name = cell(row, cols["name"])
        if not name:
            continue  # empty row
        if norm(name) in ("nom et pr", "nom et prenom", "name", "n° et nom"):
            skipped_header += 1  # repeated header row
            continue
        if name.startswith(("N°", "N", "Nom")) and len(name) < 8 and norm(name).startswith("n "):
            continue
        cls = cell(row, cols["class"]) if cols["class"] is not None else ""
        cycle_raw = cell(row, cols["cycle"]) if cols["cycle"] is not None else ""
        group_raw = cell(row, cols["group"])
        book_raw = cell(row, cols["book"]) if cols["book"] is not None else ""

        rec = build_record(name, cls, cycle_raw, group_raw, book_raw, "database", warnings)
        if rec is None:
            continue
        key = (norm(rec["name"]), norm(rec["className"]))
        if key in seen:
            warnings.append(f"Duplicate removed: {rec['name']!r} ({rec['className']})")
            continue
        seen.add(key)
        students.append(rec)

    ignored = [s for s in sheet_names if norm(s) in IGNORED_SHEETS]
    return students, warnings, target, skipped_header, ignored


def class_sort_key(c):
    order = ["cpa", "cpb", "ce1", "ce2 a", "ce2 b", "ce6 a", "ce6 b", "cm1", "cm2"]
    n = norm(c)
    if n in order:
        return (0, order.index(n))
    m = re.match(r"([a-z]+)\s*(\d+)?\s*([a-z])?", n)
    base = m.group(1) if m else n
    return (1, base)


def write_dataset(students, demo, source, warnings, extra=None):
    students.sort(key=lambda r: (r["name"].split()[0].casefold(), r["name"].casefold(), r["className"].casefold()))
    classes = sorted({r["className"] for r in students}, key=class_sort_key)
    payload = {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source": source,
        "demo": demo,
        "classes": classes,
        "students": students,
    }
    if extra:
        payload.update(extra)
    DATA_OUT.parent.mkdir(parents=True, exist_ok=True)
    DATA_OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # ---- validation report ----
    print(f"Dataset written: {DATA_OUT}")
    print(f"  students: {len(students)}  (demo={demo})")
    per_cg = {}
    for r in students:
        per_cg.setdefault(f"C{r['cycle']}-G{r['group']}", []).append(r["book"])
    for k in sorted(per_cg):
        books = sorted(set(per_cg[k]))
        print(f"  {k}: {len(per_cg[k])} students -> {', '.join(books)}")
    if warnings:
        print(f"  warnings ({len(warnings)}):")
        for w in warnings[:40]:
            print(f"    - {w}")
    if len(warnings) > 40:
        print(f"    ... and {len(warnings) - 40} more")


def demo_students():
    """Clearly-flagged demo records used ONLY until the real workbook is imported."""
    rows = [
        ("AIT FASKA Nour", "CPB", "Cycle 1", "Groupe 2", "Kid's Box 2"),
        ("ACHAMRAH Joudia", "CE1", "Cycle 1", "Groupe 3", "Kid's Box 2"),
        ("ALAOUI Sara", "CE1", "Cycle 1", "Groupe 3", "Kid's Box 2"),
        ("ALAOUI Sara", "CE2 A", "Cycle 2", "Groupe 1", "Kid's Box 3"),
        ("AMARI Anfal", "CPB", "Cycle 1", "Groupe 1", "Kid's Box 1"),
        ("AZAGANE Boutaina", "CE1", "Cycle 1", "Groupe 2", "Kid's Box 2"),
        ("BENALI Yassine", "CM1", "Cycle 3", "Groupe 1", "Kid's Box 5"),
        ("CHRAIBI Rayan", "CM2", "Cycle 3", "Groupe 2", "Kid's Box 5"),
        ("EL FASSI Adam", "CE6 A", "Cycle 2", "Groupe 2", "Kid's Box 3"),
        ("MANSOURI Aya", "CE6 B", "Cycle 2", "Groupe 3", "Kid's Box 4"),
        ("TAZI Ilyas", "CM2", "Cycle 3", "Groupe 3", "Kid's Box 6"),
    ]
    warnings = []
    students = []
    seen = set()
    for name, cls, cyc, grp, book in rows:
        rec = build_record(name, cls, cyc, grp, book, "demo", warnings)
        if rec:
            key = (norm(rec["name"]), norm(rec["className"]))
            if key not in seen:
                seen.add(key)
                students.append(rec)
    return students, warnings


def main():
    args = [a for a in sys.argv[1:]]
    if "--demo" in args or not args:
        students, warnings = demo_students()
        write_dataset(
            students,
            demo=True,
            source="DEMO DATASET - replace by running: python3 scripts/import_students.py Repartition-1.xlsx",
            warnings=warnings,
            extra={"note": "Demo records only. The real student database is loaded from the school's Repartition-1.xlsx."},
        )
        print("NOTE: writing DEMO dataset. Provide the real workbook to populate actual records.")
        return

    path = Path(args[0])
    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    students, warnings, sheet, skipped_headers, ignored = import_workbook(path)
    if not students:
        raise SystemExit("No valid student records found in the workbook.")
    write_dataset(
        students,
        demo=False,
        source=path.name,
        warnings=warnings + [f"Ignored sheets (per spec): {ignored}" if ignored else "No ignored sheets found"],
        extra={"skippedRepeatedHeaders": skipped_headers, "sheetUsed": sheet},
    )
    print("Validation: all exported records contain name, class, group and Kid's Box book.")


if __name__ == "__main__":
    main()
